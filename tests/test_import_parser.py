from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory

from django.test import SimpleTestCase
from openpyxl import Workbook

from imports.services.exceptions import CatalogImportError
from imports.services.normalizer import active, number, prepare
from imports.services.parser import MergedValues, parse_file

from .import_helpers import CLASSES, CRITICALITIES, ORIGINAL, TYPES, changed_source


class ImportParserTests(SimpleTestCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.parsed = parse_file(ORIGINAL)
        cls.plan = prepare(cls.parsed)
        cls.entries = {entry.code: entry for entry in cls.plan.levels2}

    def test_lee_hoja_correcta_y_encabezados(self):
        self.assertEqual(self.parsed.sheet, "Servicios Externos")
        self.assertEqual(len(self.parsed.rows), 97)

    def test_archivo_inexistente(self):
        with TemporaryDirectory(prefix="catalogo-import-test-") as directory:
            with self.assertRaises(CatalogImportError):
                parse_file(Path(directory) / "no-existe.xlsx")

    def test_hoja_inexistente(self):
        with changed_source(lambda book: setattr(book.active, "title", "Otra hoja")) as file:
            with self.assertRaises(CatalogImportError):
                parse_file(file)

    def test_encabezados_incompatibles(self):
        with changed_source(lambda book: setattr(book.active["A4"], "value", "OTRO")) as file:
            with self.assertRaises(CatalogImportError):
                parse_file(file)

    def test_libro_corrupto_rechazado(self):
        with TemporaryDirectory(prefix="catalogo-import-test-") as directory:
            file = Path(directory) / "no-libro.xlsx"
            file.write_bytes("datos sintéticos sin formato XLSX".encode("utf-8"))
            with self.assertRaises(CatalogImportError):
                parse_file(file)

    def test_resuelve_principal_del_rango(self):
        book = Workbook()
        sheet = book.active
        sheet["C5"] = "CODIGO"
        sheet.merge_cells("C5:C6")
        cell = MergedValues(sheet).resolve(6, 3)
        self.assertEqual(cell.value, "CODIGO")
        self.assertIsNone(cell.raw)
        self.assertEqual(cell.principal, "C5")
        self.assertEqual(cell.merged_range, "C5:C6")

    def test_no_propaga_fuera_del_rango(self):
        book = Workbook()
        sheet = book.active
        sheet["C5"] = "CODIGO"
        sheet.merge_cells("C5:C6")
        self.assertIsNone(MergedValues(sheet).resolve(7, 3).value)

    def test_combinacion_entre_atributos_rechazada(self):
        book = Workbook()
        book.active.merge_cells("C5:D5")
        with self.assertRaises(CatalogImportError):
            MergedValues(book.active)

    def test_combinacion_no_sale_del_bloque_opciones(self):
        book = Workbook()
        book.active.merge_cells("F111:F112")
        with self.assertRaises(CatalogImportError):
            MergedValues(book.active)

    def test_clasifica_continuaciones(self):
        self.assertEqual(self.plan.continuations, 49)
        self.assertEqual(sum(row.kind == "SERVICIO" for row in self.parsed.rows), 46)
        self.assertEqual(self.parsed.rows[1].kind, "CONTINUACION")

    def test_filas_42_y_67_sin_identidad(self):
        rows = {row.number: row for row in self.parsed.rows}
        for index in (42, 67):
            with self.subTest(row=index):
                self.assertEqual(rows[index].kind, "OBSERVADA")
                self.assertIsNone(rows[index].cells["C"].value)
        self.assertEqual(len(self.plan.orphan_issues), 2)

    def test_huerfana_no_aporta_atributos_al_vecino(self):
        with changed_source(lambda book: setattr(book.active["K42"], "value", 999)) as file:
            plan = prepare(parse_file(file))
            previous = next(entry for entry in plan.levels2 if entry.code == "SE.06.05")
            self.assertIsNone(previous.values["minimo"])
            orphan = next(issue for issue in plan.orphan_issues if issue.cells[0].row == 42)
            self.assertEqual(orphan.values["originales"]["K42"]["valor"], 999)

    def test_conflicto_se12_detectado(self):
        entry = next(entry for entry in self.plan.levels1 if entry.code == "SE.12")
        self.assertEqual(len(entry.issues), 1)
        self.assertEqual(entry.issues[0].kind, "CONFLICTO_N1")

    def test_nombre_canonico_se12(self):
        entries = [entry for entry in self.plan.levels1 if entry.code == "SE.12"]
        self.assertEqual(len(entries), 1)
        self.assertEqual(entries[0].values["nombre"], "Suministrar Analitica")

    def test_conserva_ambos_nombres_y_coordenadas(self):
        entry = next(entry for entry in self.plan.levels1 if entry.code == "SE.12")
        values = {cell.coordinate: cell.raw for cell in entry.cells}
        self.assertEqual(values["B99"], "Suministrar Analitica")
        self.assertEqual(values["B100"], "Mantener Tableros de Control")

    def test_se121_texto_exacto(self):
        self.assertEqual(self.entries["SE.12.1"].values["codigo"], "SE.12.1")

    def test_se122_texto_exacto(self):
        self.assertEqual(self.entries["SE.12.2"].values["codigo"], "SE.12.2")

    def test_se123_texto_exacto(self):
        self.assertEqual(self.entries["SE.12.3"].values["codigo"], "SE.12.3")

    def test_codigos_numericos_no_se_convierten_silenciosamente(self):
        with changed_source(lambda book: setattr(book.active["C99"], "value", 12.1)) as file:
            with self.assertRaises(CatalogImportError):
                prepare(parse_file(file))

    def test_inferencia_fila_101(self):
        entry = self.entries["SE.12.3"]
        self.assertEqual(entry.values["nivel1_codigo"], "SE.12")
        inference = next(
            item for item in entry.transformations if item["regla"] == "PREFIJO_EXACTO_Y_UNICO"
        )
        self.assertEqual(inference["celdas"], ["C101"])
        self.assertEqual(entry.issues[0].kind, "INFERENCIA_N1")

    def test_prefijo_ambiguo_no_se_adivina(self):
        with changed_source(lambda book: setattr(book.active["A5"], "value", "SE")) as file:
            with self.assertRaises(CatalogImportError):
                prepare(parse_file(file))

    def test_prefijo_inexistente_no_se_adivina(self):
        def change(book):
            book.active["A99"] = "OTRO"
            book.active["A100"] = "OTRO"

        with changed_source(change) as file:
            with self.assertRaises(CatalogImportError):
                prepare(parse_file(file))

    def test_atributos_null_99_a_101(self):
        for code in ("SE.12.1", "SE.12.2", "SE.12.3"):
            for field in (
                "activo", "clase", "criticidad", "tipo",
                "descripcion", "metrica", "minimo", "maximo",
            ):
                with self.subTest(code=code, field=field):
                    self.assertIsNone(self.entries[code].values[field])

    def test_activo_s(self):
        self.assertIs(active("S", "E5"), True)

    def test_activo_n(self):
        self.assertIs(active("N", "E24"), False)
        self.assertIs(self.entries["SE.05.01"].values["activo"], False)

    def test_activo_ausente(self):
        self.assertIsNone(active(None, "E99"))

    def test_activo_desconocido_rechazado(self):
        with self.assertRaises(CatalogImportError):
            active("s", "E5")

    def test_umbrales_nullable_y_cero_diferentes(self):
        self.assertIsNone(number(None, "K5"))
        self.assertEqual(number(0, "K5"), Decimal("0"))
        self.assertEqual(number("1.25", "K5"), Decimal("1.25"))

    def test_umbrales_originales_numericos(self):
        self.assertEqual(self.entries["SE.01.01"].values["minimo"], Decimal("1"))
        self.assertEqual(self.entries["SE.01.01"].values["maximo"], Decimal("100"))

    def test_umbral_no_finito_rechazado(self):
        with self.assertRaises(CatalogImportError):
            number(float("nan"), "K5")

    def test_minimo_mayor_maximo_rechazado(self):
        with changed_source(lambda book: setattr(book.active["K5"], "value", 101)) as file:
            with self.assertRaises(CatalogImportError):
                prepare(parse_file(file))

    def test_catalogos_completos_incluidos_no_usados(self):
        expected = {"CLASE": CLASSES, "CRITICIDAD": CRITICALITIES, "TIPO": TYPES}
        for entity, labels in expected.items():
            self.assertEqual({entry.code for entry in self.plan.catalogs[entity]}, labels)
        self.assertEqual(set(self.plan.active_options), {"S", "N"})

    def test_demostration_sin_correccion(self):
        entry = next(entry for entry in self.plan.catalogs["TIPO"] if entry.code == "Demostration")
        self.assertEqual(entry.values["etiqueta"], "Demostration")
        self.assertEqual(entry.cells[0].coordinate, "H113")

    def test_catalogo_no_listado_rechazado(self):
        with changed_source(lambda book: setattr(book.active["F5"], "value", "INVENTADO")) as file:
            with self.assertRaises(CatalogImportError):
                prepare(parse_file(file))

    def test_conflicto_n2_conserva_principal_y_observa(self):
        with changed_source(lambda book: setattr(book.active["F7"], "value", "RECURRENTE")) as file:
            plan = prepare(parse_file(file))
            entry = next(entry for entry in plan.levels2 if entry.code == "SE.01.01")
            self.assertEqual(entry.values["clase"], "A DEMANDA")
            self.assertEqual(entry.issues[0].kind, "CONFLICTO_N2")
            self.assertIn(
                {"celda": "F7", "principal": "F7", "valor": "RECURRENTE"},
                entry.issues[0].values["candidatos"],
            )

    def test_primera_no_vacia_dentro_identidad_si_principal_ausente(self):
        def change(book):
            book.active["I5"] = None
            book.active["I6"] = "Texto de continuación"

        with changed_source(change) as file:
            plan = prepare(parse_file(file))
            entry = next(entry for entry in plan.levels2 if entry.code == "SE.01.01")
            self.assertEqual(entry.values["descripcion"], "Texto de continuación")

    def test_texto_original_solo_dato_y_espacio_conservado(self):
        entry = self.entries["SE.01.01"]
        self.assertEqual(entry.values["descripcion"], "Revele su rollo ")

    def test_formula_no_evaluada_ni_cacheada(self):
        with changed_source(lambda book: setattr(book.active["I5"], "value", "=1+1")) as file:
            with self.assertRaises(CatalogImportError):
                parse_file(file)
