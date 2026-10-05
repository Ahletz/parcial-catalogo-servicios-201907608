"""Lectura inmutable del contrato Excel; ningún texto se ejecuta."""

from dataclasses import dataclass
from hashlib import sha256
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .exceptions import CatalogImportError

SHEET_NAME = "Servicios Externos"
HEADERS = (
    "COD.N1", "SERVICIO - Nivel 1", "COD.N2", "SERVICIO - Nivel 2",
    "ACTIVO", "CLASE DE SERVICIO", "CRITICIDAD", "TIPO DE SERVICIO",
    "Descripción", "Métrica", "Minimo", "Maximo",
)
ORIGINAL_SHA256 = "de3b478a5faeeeaebce1aa7726e0e3321188a68e41bbb656e1d17b0c5b74dcf0"


def empty(value):
    return value is None or value == ""


@dataclass(frozen=True)
class SourceFile:
    filename: str
    digest: str
    content: bytes


@dataclass(frozen=True)
class CellData:
    coordinate: str
    row: int
    column: str
    raw: object
    value: object
    principal: str
    merged_range: str | None
    data_type: str

    def evidence(self):
        def json_value(value):
            return value.isoformat() if hasattr(value, "isoformat") else value

        return {
            "valor": json_value(self.raw),
            "interpretado_por_combinacion": json_value(self.value),
            "principal": self.principal,
            "rango": self.merged_range,
            "tipo": self.data_type,
        }


@dataclass(frozen=True)
class SourceRow:
    number: int
    kind: str
    cells: dict[str, CellData]


@dataclass(frozen=True)
class ParsedWorkbook:
    source: SourceFile
    sheet: str
    rows: list[SourceRow]
    options: dict[str, list[CellData]]
    merged_ranges: list[str]


class MergedValues:
    """Devuelve únicamente la principal del rango que contiene la celda."""

    def __init__(self, sheet):
        self.sheet = sheet
        self.ranges = list(sheet.merged_cells.ranges)
        self.index = {}
        for area in self.ranges:
            in_catalog = (
                area.min_col <= 12 and area.max_col >= 1
                and area.min_row <= 101 and area.max_row >= 5
            )
            in_options = (
                area.min_col <= 8 and area.max_col >= 5
                and area.min_row <= 122 and area.max_row >= 112
            )
            relevant = in_catalog or in_options
            if relevant and area.min_col != area.max_col:
                raise CatalogImportError(
                    f"Combinación entre atributos incompatible: {area}."
                )
            if relevant and area.min_row < 5:
                raise CatalogImportError(f"Combinación cruza los encabezados: {area}.")
            if relevant and area.min_row <= 101 < area.max_row:
                raise CatalogImportError(f"Combinación cruza el catálogo: {area}.")
            if in_options and (area.min_row < 112 or area.max_row > 122):
                raise CatalogImportError(f"Combinación cruza las opciones: {area}.")
            if relevant:
                for row in range(area.min_row, area.max_row + 1):
                    self.index[(row, area.min_col)] = area

    def resolve(self, row, column):
        cell = self.sheet.cell(row, column)
        principal = cell
        merged_range = None
        area = self.index.get((row, column))
        if area is not None:
            principal = self.sheet.cell(area.min_row, area.min_col)
            merged_range = str(area)
        if cell.data_type == "f" or principal.data_type == "f":
            raise CatalogImportError(
                f"Fórmula no admitida en el contrato de datos: {principal.coordinate}."
            )
        return CellData(
            coordinate=cell.coordinate, row=row,
            column=get_column_letter(column), raw=cell.value,
            value=principal.value, principal=principal.coordinate,
            merged_range=merged_range, data_type=principal.data_type,
        )


def read_source(file):
    path = Path(file)
    if not path.is_file():
        raise CatalogImportError("El archivo de importación no existe o no es un archivo.")
    try:
        content = path.read_bytes()
    except OSError as exc:
        raise CatalogImportError("No se pudo leer el archivo de importación.") from exc
    return SourceFile(path.name, sha256(content).hexdigest(), content)


def parse_workbook(source):
    try:
        book = load_workbook(
            BytesIO(source.content), data_only=False, read_only=False,
            keep_links=False,
        )
    except Exception as exc:
        raise CatalogImportError("El archivo no es un libro XLSX legible.") from exc
    try:
        if SHEET_NAME not in book.sheetnames:
            raise CatalogImportError(f"No existe la hoja requerida: {SHEET_NAME}.")
        sheet = book[SHEET_NAME]
        if tuple(sheet.cell(4, col).value for col in range(1, 13)) != HEADERS:
            raise CatalogImportError("Encabezados A4:L4 incompatibles con el contrato.")
        merged = MergedValues(sheet)
        rows = []
        for number in range(5, 102):
            cells = {
                get_column_letter(col): merged.resolve(number, col)
                for col in range(1, 13)
            }
            identity = cells["C"]
            if not empty(identity.raw):
                kind = "SERVICIO"
            elif not empty(identity.value) and identity.merged_range:
                kind = "CONTINUACION"
            elif any(not empty(cell.value) for cell in cells.values()):
                kind = "OBSERVADA"
            else:
                kind = "VACIA"
            rows.append(SourceRow(number, kind, cells))
        options = {
            get_column_letter(col): [
                merged.resolve(row, col) for row in range(112, 123)
            ]
            for col in range(5, 9)
        }
        return ParsedWorkbook(
            source, SHEET_NAME, rows, options,
            sorted(str(area) for area in merged.ranges),
        )
    finally:
        book.close()


def parse_file(file):
    return parse_workbook(read_source(file))
